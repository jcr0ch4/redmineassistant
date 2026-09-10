"""Download das atividades ativas atribuídas ao usuário e geração da planilha de atuação.

Gera o arquivo atividades_ativas.xlsx com as colunas editáveis (Status, % Concluído,
Data Prevista, Horas Apontadas, Comentário). Serve de template para o script de
atualização do Redmine.

Uso:
    python download_atividades.py
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from redmine_api import STATUS_ATIVOS, criar_api_da_config

BASE_DIR = Path(__file__).resolve().parent
ARQUIVO_SAIDA = BASE_DIR / "atividades_ativas.xlsx"

COLUNAS = [
    "ID",
    "Projeto",
    "Tipo",
    "Status",
    "Prioridade",
    "Assunto",
    "Data Início",
    "Data Prevista",
    "% Concluído",
    "Horas Estimadas",
    "Horas Gastas (Redmine)",
    "Horas Apontadas",
    "Comentário",
]

# Mapeamento nome do status -> cor de destaque
CORES_STATUS = {
    "Nova": "FFC7CE",
    "Em andamento": "FFEB9C",
    "Especificação": "C6EFCE",
    "Backlog": "E2EFDA",
}


def gerar_planilha(caminho: Path):
    api = criar_api_da_config()
    usuario = api.get_usuario_atual()
    issues = api.get_issues_ativas()

    wb = Workbook()

    # ---------- Aba Ativos ----------
    ws = wb.active
    ws.title = "Ativos"

    estilo_header = Font(bold=True, color="FFFFFF")
    fill_header = PatternFill("solid", fgColor="4472C4")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    borda = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col, titulo in enumerate(COLUNAS, start=1):
        cel = ws.cell(row=1, column=col, value=titulo)
        cel.font = estilo_header
        cel.fill = fill_header
        cel.alignment = align_center
        cel.border = borda

    for linha, issue in enumerate(issues, start=2):
        valores = [
            issue.get("id", ""),
            (issue.get("project") or {}).get("name", ""),
            (issue.get("tracker") or {}).get("name", ""),
            (issue.get("status") or {}).get("name", ""),
            (issue.get("priority") or {}).get("name", ""),
            issue.get("subject", ""),
            issue.get("start_date", ""),
            issue.get("due_date", "") or "",
            issue.get("done_ratio", 0),
            issue.get("estimated_hours", ""),
            issue.get("spent_hours", 0),
            "",  # Horas Apontadas (a preencher pelo usuário)
            "",  # Comentário
        ]
        for col, valor in enumerate(valores, start=1):
            cel = ws.cell(row=linha, column=col, value=valor)
            cel.border = borda
            cel.alignment = Alignment(vertical="center", wrap_text=(col == 6))
            if col == 9:
                cel.number_format = '0"%"'

        status = (issue.get("status") or {}).get("name")
        if status in CORES_STATUS:
            preenche = PatternFill("solid", fgColor=CORES_STATUS[status])
            ws.cell(row=linha, column=4).fill = preenche

    larguras = [8, 32, 18, 16, 11, 60, 12, 12, 11, 13, 13, 13, 40]
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura

    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{len(issues)+1}"
    ws.freeze_panes = "A2"

    # ---------- Aba Config (validação) ----------
    ws2 = wb.create_sheet("Config")

    ws2.cell(row=1, column=1, value="Status disponíveis no Redmine")
    ws2.cell(row=1, column=1).font = Font(bold=True)
    for i, nome in enumerate(sorted(STATUS_ATIVOS)):
        ws2.cell(row=i + 2, column=1, value=nome)

    # Validação de lista para a coluna Status (D)
    if len(issues) > 0:
        from openpyxl.worksheet.datavalidation import DataValidation
        lista_status = f"Config!$A$2:$A${len(STATUS_ATIVOS)+1}"
        dv = DataValidation(type="list", formula1=lista_status, allow_blank=True)
        dv.error = "Escolha um status da lista (Nova, Em andamento, Backlog, Especificação)."
        dv.errorTitle = "Status inválido"
        ws.add_data_validation(dv)
        dv.add(f"D2:D{len(issues)+1}")

    # Registro de metadados (usado pelo script de atualização)
    ws2.cell(row=9, column=1, value="usuario")
    ws2.cell(row=9, column=2, value=usuario.get("login", ""))
    ws2.cell(row=10, column=1, value="usuario_id")
    ws2.cell(row=10, column=2, value=usuario.get("id", ""))
    ws2.cell(row=11, column=1, value="gerado_em")
    ws2.cell(row=11, column=2, value=datetime.now().isoformat(timespec="seconds"))

    wb.save(caminho)
    print(f"Planilha de atuação gerada: {caminho}")
    print(f"Total de atividades ativas: {len(issues)}")
    print("Edite a planilha (Status, % Concluído, Data Prevista, Horas Apontadas, Comentário) e depois rode atualizar_redmine.py.")


if __name__ == "__main__":
    gerar_planilha(ARQUIVO_SAIDA)
